import openpyxl
import json
import re
import os

EXCEL_PATH = r"C:\Users\Elif\Downloads\ESOGÜİF_2026-2027-GÜZ-DÖNEMİ-DERS-PROGRAMI.xlsx"
OUTPUT_JSON = r"C:\Users\Elif\.gemini\antigravity\scratch\esogu-ders-programi\data\schedule_data.json"

wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
ws = wb["DERS PROGRAMI"]

current_section_title = ""
current_category = "ZORUNLU"

raw_rows = []

for r in range(1, ws.max_row + 1):
    c1 = ws.cell(r, 1).value
    c2 = ws.cell(r, 2).value
    c3 = ws.cell(r, 3).value
    c4 = ws.cell(r, 4).value
    c5 = ws.cell(r, 5).value
    c6 = ws.cell(r, 6).value
    c7 = ws.cell(r, 7).value
    c8 = ws.cell(r, 8).value
    c9 = ws.cell(r, 9).value

    if c1 and not c4:
        s = str(c1).strip()
        if "PROGRAMI" in s and "ESOG" in s:
            continue
        if "SINIF" in s:
            current_section_title = s
            current_category = "ZORUNLU"
        elif "SEÇMELİ" in s:
            current_category = "SEÇMELİ"
        continue

    if c1 and str(c1).strip() == "Dersin Kodu":
        continue

    if c1 and c3 and c4:
        raw_rows.append({
            "section_title": current_section_title,
            "category": current_category,
            "code": str(c1).strip(),
            "sube": str(c2 or "").strip(),
            "name": str(c3).strip(),
            "time_str": str(c4).strip(),
            "teo": int(c5) if c5 is not None and str(c5).isdigit() else 0,
            "uyg": int(c6) if c6 is not None and str(c6).isdigit() else 0,
            "derslik": str(c7 or "").strip(),
            "unvan": str(c8 or "").strip(),
            "hoca": str(c9 or "").strip()
        })

print(f"Parsed {len(raw_rows)} raw rows from Excel.")

# Formation Course definitions
FORMATION_KEYWORDS = [
    "EĞİTİME GİRİŞ",
    "ÖĞRETİM İLKE VE YÖNTEMLERİ",
    "EĞİTİMDE ÖLÇME VE DEĞERLENDİRME",
    "ÖĞRETİM TEKNOLOJİLERİ",
    "ÖZEL ÖĞRETİM YÖNTEMLERİ",
    "ÖĞRETMENLİK UYGULAMASI I"
]

# ECTS and Credit mapping from official 2024 / 2021 curriculum
OFFICIAL_AKTS_MAP = {
    # 1. Sınıf
    "181011001": 2, # ATATÜRK İLKE.VE İNK.TARİHİ I
    "181011002": 2, # İNGİLİZCE I
    "181011005": 3, # İNGİLİZCE I
    "181011004": 2, # TÜRK DİLİ I
    "181111024": 3, # KURAN OKUMA VE TECVİD I
    "181111025": 3, # İSLAM İNANÇ ESASLARI
    "181111026": 3, # ARAP DİLİ VE BELAGATI
    "181111038": 4, # MANTIK
    "181111039": 3, # SİYER
    "181111040": 5, # TEFSİR TARİHİ VE USULÜ
    "181111041": 4, # EĞİTİME GİRİŞ (PF)
    # 2. Sınıf
    "181113035": 4, # HADİS I
    "181113036": 4, # İSLAM TARİHİ I
    "181113037": 3, # KUR'AN OKUMA VE TECVİD III
    "181113038": 3, # OSMANLI TÜRKÇESİ
    "181113039": 3, # PSİKOLOJİYE GİRİŞ
    "181113040": 3, # SOSYOLOJİYE GİRİŞ
    "181113041": 4, # TEFSİR I
    "181113050": 4, # ÖĞRETİM İLKE VE YÖNTEMLERİ (PF)
    # 3. Sınıf
    "181115060": 3, # DİN EĞİTİMİ
    "181115061": 5, # DİNLER TARİHİ
    "181115062": 5, # İSLAM HUKUKU I
    "181115063": 3, # KUR'AN OKUMA VE TECVİD V
    "181115064": 5, # SİSTEMATİK KELAM I
    "181115065": 4, # EĞİTİMDE ÖLÇME VE DEĞERLENDİRME (PF)
    "181115066": 3, # ÖĞRETİM TEKNOLOJİLERİ (PF)
    # 4. Sınıf (2021 P.F. - Kredi bazlı, ama AKTS karşılıkları)
    "181117038": 4, # KURAN OKUMA VE TECVİD VII
    "181117045": 4, # HADİS III
    "181117046": 3, # TÜRK İSLAM SANATLARI TARİHİ
    "181117047": 4, # SİSTEMATİK KELAM III
    "181117048": 3, # DİN EĞİTİMİ
    "181117049": 2, # TÜRK DİN MUSİKİSİ(NAZARİYAT)
    "181117050": 4, # DİN FELSEFESİ I
    "181117051": 3, # DİNİ HİTABET VE MESLEKİ UYGULAMA
    "181117063": 3, # ÖZEL ÖĞRETİM YÖNTEMLERİ (PF)
    "181117069": 5, # ÖĞRETMENLİK UYGULAMASI I (PF)
}

# Group raw rows into Course Offerings (Ders + Şube)
grouped = {}

for r in raw_rows:
    key = (r["section_title"], r["category"], r["code"], r["sube"])
    if key not in grouped:
        # Determine Grade (Sınıf)
        st = r["section_title"]
        grade = 1
        if "İKİNCİ SINIF" in st:
            grade = 2
        elif "ÜÇÜNCÜ SINIF" in st:
            grade = 3
        elif "DÖRDÜNCÜ SINIF" in st:
            grade = 4

        # Determine Teaching Type (1. Öğretim vs 2. Öğretim)
        teaching_type = "1. Öğretim"
        if "İKİNCİ ÖĞRETİM" in st or "İKİNCİ ÖĞRETİM" in r["name"]:
            teaching_type = "2. Öğretim"
        elif r["name"] == "DİNLER TARİHİ PERSPEKTİFİNDEN HAÇLI SEFERLERİ" and r["sube"] == "B":
            # Section B opened from 1st teaching for 2nd teaching students
            teaching_type = "2. Öğretim"

        # Determine Curriculum Template
        is_2021_pf = False
        if "2021 P.F." in st or "2021 P.F." in r["name"]:
            is_2021_pf = True

        # Determine Formation Course
        is_formation = False
        upper_name = r["name"].upper()
        for kw in FORMATION_KEYWORDS:
            if kw in upper_name:
                is_formation = True
                break

        # Calculate Kredi
        kredi = r["teo"] + (r["uyg"] // 2)
        if kredi == 0:
            kredi = r["teo"] if r["teo"] > 0 else 2

        # Calculate AKTS
        akts = OFFICIAL_AKTS_MAP.get(r["code"], 3) # default 3 for electives
        if is_formation and akts == 3:
            akts = 3

        # Gender restrictions: Yalnızca zorunlu Kur'an Okuma ve Tecvid derslerinde şubelere göre ayrım vardır.
        # Diğer tüm derslerde (Dini Hitabet, Zorunlu dersler ve Seçmeliler) şube cinsiyet ayrımı KESİNLİKLE YOKTUR!
        gender_restriction = "HEPSİ"
        if r["category"] == "ZORUNLU" and ("KURAN OKUMA" in upper_name or "KUR'AN OKUMA" in upper_name):
            # 2021 P.F. eski kodlu tek şube olanlar ortak açılır veya her iki cinsiyete açıktır
            if "2021 P.F." in r["name"]:
                gender_restriction = "HEPSİ"
            elif grade == 3 and teaching_type == "2. Öğretim":
                if r["sube"] in ["A", "B"]:
                    gender_restriction = "KIZ"
                elif r["sube"] == "C":
                    gender_restriction = "ERKEK"
            else:
                if r["sube"] in ["A", "B", "C"]:
                    gender_restriction = "KIZ"
                elif r["sube"] in ["D", "E"]:
                    gender_restriction = "ERKEK"

        # Mantık restriction
        mantik_restriction = "HEPSİ"
        if r["code"] == "181111038" or "MANTIK" in upper_name:
            if r["sube"] == "A" and teaching_type == "1. Öğretim":
                mantik_restriction = "İLK_DEFA"
            else:
                mantik_restriction = "ALTTAN"

        grouped[key] = {
            "id": f"{r['code']}_{r['sube']}_{grade}_{'2O' if teaching_type == '2. Öğretim' else '1O'}",
            "code": r["code"],
            "sube": r["sube"],
            "name": r["name"],
            "section_title": r["section_title"],
            "category": r["category"], # ZORUNLU or SEÇMELİ
            "grade": grade,
            "teaching_type": teaching_type,
            "is_2021_pf": is_2021_pf,
            "curriculum": "2021 P.F." if is_2021_pf else "AKTS 2024",
            "is_formation": is_formation,
            "gender_restriction": gender_restriction,
            "mantik_restriction": mantik_restriction,
            "kredi": kredi,
            "akts": akts,
            "teo": r["teo"],
            "uyg": r["uyg"],
            "derslik": r["derslik"],
            "unvan": r["unvan"],
            "hoca": r["hoca"],
            "is_online": (r["derslik"].lower() == "online"),
            "has_internship": False,
            "time_slots": []
        }

    # Parse time slot
    t_str = r["time_str"]
    if t_str == "OKULDA UYGULAMA":
        grouped[key]["has_internship"] = True
    else:
        m = re.match(r"^(\d+):0\s+(PAZARTESİ|SALI|ÇARŞAMBA|PERŞEMBE|CUMA)$", t_str)
        if m:
            hour = int(m.group(1))
            day = m.group(2)
            grouped[key]["time_slots"].append({
                "day": day,
                "hour": hour,
                "raw": t_str,
                "label": f"{hour:02d}:00 - {hour+1:02d}:00"
            })
        else:
            print(f"Warning: unrecognized time format '{t_str}' in {r['name']}")

course_list = list(grouped.values())
print(f"Total unique course-section offerings: {len(course_list)}")

# Count statistics
formation_count = sum(1 for c in course_list if c["is_formation"])
io_count = sum(1 for c in course_list if c["teaching_type"] == "2. Öğretim")
pf_count = sum(1 for c in course_list if c["is_2021_pf"])
print(f"Formation courses: {formation_count}")
print(f"2. Öğretim courses: {io_count}")
print(f"2021 P.F. courses: {pf_count}")

# Save JSON & JS
OUTPUT_JS = os.path.join(os.path.dirname(OUTPUT_JSON), "schedule_data.js")
os.makedirs(os.path.dirname(OUTPUT_JSON), exist_ok=True)

with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
    json.dump(course_list, f, ensure_ascii=False, indent=2)

with open(OUTPUT_JS, "w", encoding="utf-8") as f:
    f.write("window.SCHEDULE_DATA = " + json.dumps(course_list, ensure_ascii=False, indent=2) + ";\n")

print(f"Saved successfully to {OUTPUT_JSON} and {OUTPUT_JS}")
